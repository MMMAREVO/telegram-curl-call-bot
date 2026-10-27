import csv
import io

from openpyxl import Workbook, load_workbook

from . import excel_store
from .campaigns import Campaign, VariableList
from .phone import validate_phone

# Excel/CSV-файлы, которыми админ обменивается с ботом: пользователи, тест-планы,
# списки переменных кампаний.

_TRUE_MARKS = {"да", "1", "true", "yes", "y", "x", "+", "админ", "admin"}
# Сколько проблемных строк перечислять в ответе, чтобы не упереться в лимит сообщения
_MAX_ERRORS = 15


def _xlsx_bytes(header: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.append(header)
    for row in rows:
        excel_store.append_row(sheet, row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(value) -> str:
    """Excel хранит номера и коды числами (998901234567.0) — приводим к строке без ".0"."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_table(buf: io.BytesIO, filename: str) -> tuple[list[str], list[list[str]], str | None]:
    """(заголовок, строки, ошибка). Понимает .xlsx и .csv; пустые строки пропускает."""
    if filename.lower().endswith(".csv"):
        raw = buf.getvalue()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        table = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        try:
            sheet = load_workbook(buf, data_only=True).active
        except Exception as exc:
            return [], [], f"Не удалось открыть файл как Excel: {exc}"
        table = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]

    table = [row for row in table if any(row)]
    if not table:
        return [], [], "Файл пустой."
    header, *rows = table
    while header and not header[-1]:  # пустые колонки справа (бывают от форматирования в Excel)
        header.pop()
    width = len(header)
    return header, [(row + [""] * width)[:width] for row in rows], None


def _cap(errors: list[str]) -> list[str]:
    if len(errors) > _MAX_ERRORS:
        return errors[:_MAX_ERRORS] + [f"…и ещё {len(errors) - _MAX_ERRORS}"]
    return errors


# --- Пользователи ---

def build_users_file(users: list[dict]) -> bytes:
    return _xlsx_bytes(
        excel_store.USERS_HEADER,
        [
            [u["telegram_id"], u["name"], u["phone"], u["registered_at"], excel_store.ADMIN_MARK if u["is_admin"] else ""]
            for u in users
        ],
    )


def parse_users_file(buf: io.BytesIO, filename: str) -> tuple[list[dict], list[str]]:
    """Разбирает исправленную админом таблицу пользователей. Возвращает (пользователи, ошибки)."""
    header, rows, error = _read_table(buf, filename)
    if error:
        return [], [error]
    col = {name: i for i, name in enumerate(header) if name}
    missing = [c for c in ("telegram_id", "name", "phone") if c not in col]
    if missing:
        return [], [f"В файле не хватает колонок: {', '.join(missing)}"]

    users: list[dict] = []
    errors: list[str] = []
    seen: set[int] = set()
    for line, row in enumerate(rows, start=2):
        try:
            telegram_id = int(row[col["telegram_id"]])
        except ValueError:
            errors.append(f"Строка {line}: telegram_id должен быть числом")
            continue
        if telegram_id in seen:
            errors.append(f"Строка {line}: telegram_id {telegram_id} повторяется")
            continue
        seen.add(telegram_id)
        name = row[col["name"]]
        if not name:
            errors.append(f"Строка {line}: пустое имя")
            continue
        phone, phone_error = validate_phone(row[col["phone"]])
        if phone_error:
            errors.append(f"Строка {line}: {phone_error}")
            continue
        users.append({
            "telegram_id": telegram_id,
            "name": name,
            "phone": phone,
            "registered_at": row[col["registered_at"]] if "registered_at" in col else "",
            "is_admin": "is_admin" in col and row[col["is_admin"]].lower() in _TRUE_MARKS,
        })
    if not users and not errors:
        errors.append("В файле нет ни одного пользователя.")
    return users, _cap(errors)


# --- Тест-планы ---

def build_testcases_file(rows: list[dict]) -> bytes:
    header = excel_store.TESTCASES_HEADER
    return _xlsx_bytes(header, [[r.get(c, "") for c in header] for r in rows])


# --- Списки переменных ---

def build_var_list_file(var_list: VariableList) -> bytes:
    return _xlsx_bytes(var_list.columns, var_list.rows)


def parse_var_list_file(buf: io.BytesIO, filename: str, campaign: Campaign) -> tuple[VariableList | None, list[str]]:
    """Первая строка — имена переменных, остальные — наборы значений. Колонки должны
    покрывать все {{переменные}} курла кампании; лишние колонки сохраняются, но не используются."""
    header, rows, error = _read_table(buf, filename)
    if error:
        return None, [error]
    if len(set(header)) != len(header) or not all(header):
        return None, ["В первой строке должны быть имена переменных — без пустых и повторов."]
    missing = [v for v in campaign.variables if v not in header]
    if missing:
        return None, [
            f"Нет колонок для переменных курла: {', '.join(missing)}. "
            f"Колонки в файле: {', '.join(header)}"
        ]
    if not rows:
        return None, ["В файле нет ни одной строки со значениями."]

    used = [header.index(v) for v in campaign.variables]
    errors = [
        f"Строка {line}: не заполнено {', '.join(header[i] for i in used if not row[i])}"
        for line, row in enumerate(rows, start=2)
        if any(not row[i] for i in used)
    ]
    if errors:
        return None, _cap(errors)
    return VariableList(columns=header, rows=rows), []
