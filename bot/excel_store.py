import asyncio
import datetime as dt
import time
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import config

_lock = asyncio.Lock()
_SAVE_RETRY_ATTEMPTS = 5
_SAVE_RETRY_DELAY = 1.0


def _save_with_retry(wb: Workbook) -> None:
    """На Windows открытый в Excel файл блокируется на запись — пробуем ещё раз
    несколько раз, прежде чем сдаться."""
    last_exc: Optional[PermissionError] = None
    for _ in range(_SAVE_RETRY_ATTEMPTS):
        try:
            wb.save(config.EXCEL_PATH)
            return
        except PermissionError as exc:
            last_exc = exc
            time.sleep(_SAVE_RETRY_DELAY)
    raise PermissionError(
        f"файл {config.EXCEL_PATH} занят другим процессом (возможно, открыт в Excel)"
    ) from last_exc

# Строки с такими первыми символами Excel воспринял бы как формулу (например, имя
# "=HYPERLINK(...)" выполнилось бы у админа при открытии файла) — пишем их строго как текст
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _as_text_if_formula(cell, value) -> None:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"


def set_cell(sheet: Worksheet, row: int, column: int, value) -> None:
    _as_text_if_formula(sheet.cell(row=row, column=column, value=value), value)


def append_row(sheet: Worksheet, values: list) -> None:
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        _as_text_if_formula(cell, value)


USERS_HEADER = ["telegram_id", "name", "phone", "registered_at", "is_admin"]
# Так в колонке is_admin помечается админ, назначенный из бота
ADMIN_MARK = "да"
CALLS_HEADER = [
    "timestamp", "telegram_id", "name", "phone", "campaign",
    "params", "status", "error", "sid", "call_result", "log", "log_file",
    "transcription", "transcript_file", "comment", "call_id", "round", "test_case", "test_plan",
    "audio_file",
    # Больше не заполняются (замер по логу студии убран), но колонки остаются: порядок
    # CALLS_HEADER должен совпадать с уже существующим заголовком листа
    "latency_avg_ms", "latency_median_ms", "latency_min_ms", "latency_max_ms", "latency_turns_ms",
    # Пауза для собеседника по логам VictoriaLogs (см. pause.py), мс
    "pause_avg_ms", "pause_median_ms", "pause_min_ms", "pause_max_ms", "pause_turns_ms",
    # Пауза "как слышится" — по стерео-записи до первого звука бота (см. audio_pause.py), мс
    "heard_avg_ms", "heard_median_ms", "heard_min_ms", "heard_max_ms", "heard_turns_ms",
]
# Один раунд тестирования на сценарий: campaign_key -> текущий номер раунда.
ROUNDS_HEADER = ["campaign_key", "current_round", "updated_at", "updated_by"]
# Тест-кейсы, загружаемые админом: несколько строк на один test_plan.
TESTCASES_HEADER = ["test_plan", "order", "description", "campaign"]
# Прогресс исполнителя по тест-плану: индекс текущего тест-кейса.
TESTPROGRESS_HEADER = ["telegram_id", "test_plan", "current_index", "updated_at"]


def _migrate_header(sheet: Worksheet, full_header: list[str]) -> None:
    """Дописывает недостающие колонки в конец существующего заголовка, никогда не
    переставляя и не удаляя то, что там уже есть — иначе разъедутся старые строки."""
    existing = {sheet.cell(row=1, column=i).value for i in range(1, sheet.max_column + 1)}
    next_col = sheet.max_column + 1
    for header in full_header:
        if header not in existing:
            sheet.cell(row=1, column=next_col, value=header)
            next_col += 1


def _open_workbook() -> Workbook:
    path = Path(config.EXCEL_PATH)
    if path.exists():
        wb = load_workbook(path)
        _migrate_header(wb["Calls"], CALLS_HEADER)
        _migrate_header(wb["Users"], USERS_HEADER)
        if "Rounds" not in wb.sheetnames:
            rounds_sheet = wb.create_sheet("Rounds")
            rounds_sheet.append(ROUNDS_HEADER)
        if "TestCases" not in wb.sheetnames:
            tc_sheet = wb.create_sheet("TestCases")
            tc_sheet.append(TESTCASES_HEADER)
        if "TestProgress" not in wb.sheetnames:
            tp_sheet = wb.create_sheet("TestProgress")
            tp_sheet.append(TESTPROGRESS_HEADER)
        return wb
    wb = Workbook()
    users_sheet = wb.active
    users_sheet.title = "Users"
    users_sheet.append(USERS_HEADER)
    calls_sheet = wb.create_sheet("Calls")
    calls_sheet.append(CALLS_HEADER)
    rounds_sheet = wb.create_sheet("Rounds")
    rounds_sheet.append(ROUNDS_HEADER)
    tc_sheet = wb.create_sheet("TestCases")
    tc_sheet.append(TESTCASES_HEADER)
    tp_sheet = wb.create_sheet("TestProgress")
    tp_sheet.append(TESTPROGRESS_HEADER)
    _save_with_retry(wb)
    return wb


def _find_user_row(sheet: Worksheet, telegram_id: int) -> Optional[int]:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == telegram_id:
            return row
    return None


def _list_user_ids_sync() -> list[int]:
    sheet = _open_workbook()["Users"]
    return [v for v in (sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)) if v is not None]


async def list_user_ids() -> list[int]:
    async with _lock:
        return await asyncio.to_thread(_list_user_ids_sync)


def _get_user_sync(telegram_id: int) -> Optional[dict]:
    wb = _open_workbook()
    sheet = wb["Users"]
    row = _find_user_row(sheet, telegram_id)
    if row is None:
        return None
    return {
        "name": sheet.cell(row=row, column=2).value,
        "phone": str(sheet.cell(row=row, column=3).value),
    }


async def get_user(telegram_id: int) -> Optional[dict]:
    async with _lock:
        return await asyncio.to_thread(_get_user_sync, telegram_id)


def _register_user_sync(telegram_id: int, name: str, phone: str) -> None:
    wb = _open_workbook()
    sheet = wb["Users"]
    row = _find_user_row(sheet, telegram_id)
    now = dt.datetime.now().isoformat(timespec="seconds")
    if row is None:
        append_row(sheet, [telegram_id, name, phone, now])
    else:
        set_cell(sheet, row, 2, name)
        set_cell(sheet, row, 3, phone)
    _save_with_retry(wb)


async def register_user(telegram_id: int, name: str, phone: str) -> None:
    async with _lock:
        await asyncio.to_thread(_register_user_sync, telegram_id, name, phone)


def _user_record(sheet: Worksheet, row: int) -> dict:
    record = {h: sheet.cell(row=row, column=i + 1).value for i, h in enumerate(USERS_HEADER)}
    record["phone"] = str(record["phone"] or "")
    record["is_admin"] = bool(record["is_admin"])
    return record


def _list_users_sync() -> list[dict]:
    sheet = _open_workbook()["Users"]
    return [
        _user_record(sheet, row) for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=1).value is not None
    ]


async def list_users() -> list[dict]:
    """Все пользователи: telegram_id, name, phone, registered_at, is_admin (bool)."""
    async with _lock:
        return await asyncio.to_thread(_list_users_sync)


def _update_user_sync(telegram_id: int, fields: dict) -> bool:
    wb = _open_workbook()
    sheet = wb["Users"]
    row = _find_user_row(sheet, telegram_id)
    if row is None:
        return False
    for key, value in fields.items():
        if key == "is_admin":
            value = ADMIN_MARK if value else ""
        set_cell(sheet, row, USERS_HEADER.index(key) + 1, value)
    _save_with_retry(wb)
    return True


async def update_user(telegram_id: int, **fields) -> bool:
    """Меняет name / phone / is_admin пользователя. False — такого пользователя нет."""
    async with _lock:
        return await asyncio.to_thread(_update_user_sync, telegram_id, fields)


def _delete_user_sync(telegram_id: int) -> bool:
    wb = _open_workbook()
    removed = _delete_rows_where(wb["Users"], lambda v: v[0] == telegram_id)
    _save_with_retry(wb)
    return removed > 0


async def delete_user(telegram_id: int) -> bool:
    async with _lock:
        return await asyncio.to_thread(_delete_user_sync, telegram_id)


def _replace_users_sync(users: list[dict]) -> int:
    wb = _open_workbook()
    sheet = wb["Users"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for user in users:
        values = {**user, "is_admin": ADMIN_MARK if user.get("is_admin") else ""}
        append_row(sheet, [values.get(col, "") for col in USERS_HEADER])
    _save_with_retry(wb)
    return len(users)


async def replace_users(users: list[dict]) -> int:
    """Полностью заменяет таблицу пользователей (загрузка исправленного Excel админом).
    Возвращает, сколько пользователей записано."""
    async with _lock:
        return await asyncio.to_thread(_replace_users_sync, users)


def _add_call_row_sync(fields: dict) -> int:
    wb = _open_workbook()
    sheet = wb["Calls"]
    append_row(sheet, [fields.get(col, "") for col in CALLS_HEADER])
    row_index = sheet.max_row
    _save_with_retry(wb)
    return row_index


async def add_call_row(**fields) -> int:
    async with _lock:
        return await asyncio.to_thread(_add_call_row_sync, fields)


def _update_call_row_sync(row_index: int, fields: dict) -> None:
    wb = _open_workbook()
    sheet = wb["Calls"]
    for key, value in fields.items():
        col = CALLS_HEADER.index(key) + 1
        set_cell(sheet, row_index, col, value)
    _save_with_retry(wb)


async def update_call_row(row_index: int, **fields) -> None:
    async with _lock:
        await asyncio.to_thread(_update_call_row_sync, row_index, fields)


def _find_round_row(sheet: Worksheet, campaign_key: str) -> Optional[int]:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == campaign_key:
            return row
    return None


def _get_current_round_sync(campaign_key: str) -> Optional[int]:
    wb = _open_workbook()
    sheet = wb["Rounds"]
    row = _find_round_row(sheet, campaign_key)
    if row is None:
        return None
    return sheet.cell(row=row, column=2).value


async def get_current_round(campaign_key: str) -> Optional[int]:
    async with _lock:
        return await asyncio.to_thread(_get_current_round_sync, campaign_key)


def _create_new_round_sync(campaign_key: str, admin_telegram_id: int) -> int:
    wb = _open_workbook()
    sheet = wb["Rounds"]
    row = _find_round_row(sheet, campaign_key)
    current = sheet.cell(row=row, column=2).value if row is not None else 0
    new_round = (current or 0) + 1
    now = dt.datetime.now().isoformat(timespec="seconds")
    if row is None:
        append_row(sheet, [campaign_key, new_round, now, admin_telegram_id])
    else:
        set_cell(sheet, row, 2, new_round)
        set_cell(sheet, row, 3, now)
        set_cell(sheet, row, 4, admin_telegram_id)
    _save_with_retry(wb)
    return new_round


async def create_new_round(campaign_key: str, admin_telegram_id: int) -> int:
    async with _lock:
        return await asyncio.to_thread(_create_new_round_sync, campaign_key, admin_telegram_id)


def _ensure_round_sync(campaign_key: str) -> int:
    wb = _open_workbook()
    sheet = wb["Rounds"]
    row = _find_round_row(sheet, campaign_key)
    if row is not None:
        return sheet.cell(row=row, column=2).value
    now = dt.datetime.now().isoformat(timespec="seconds")
    append_row(sheet, [campaign_key, 1, now, "auto"])
    _save_with_retry(wb)
    return 1


async def ensure_round(campaign_key: str) -> int:
    """Возвращает текущий раунд для сценария, автоматически создавая раунд 1,
    если админ ещё не завёл ни одного (чтобы не блокировать тестировщиков)."""
    async with _lock:
        return await asyncio.to_thread(_ensure_round_sync, campaign_key)


def is_error_row(record: dict) -> bool:
    """Звонок считается ошибочным, если сам curl не прошёл, либо тестировщик
    явно пометил его кнопкой 'Ошибка. Есть комментарий'."""
    if record.get("status") == "error":
        return True
    comment = (record.get("comment") or "").strip()
    return bool(comment) and comment != "Успешный звонок"


def _query_calls_sync(
    campaign_title: str, telegram_id: Optional[int], only_errors: bool, round_filter: Optional[int]
) -> list[dict]:
    wb = _open_workbook()
    sheet = wb["Calls"]
    rows = []
    for r in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=r, column=i + 1).value for i in range(len(CALLS_HEADER))]
        record = dict(zip(CALLS_HEADER, values))
        if record.get("campaign") != campaign_title:
            continue
        if telegram_id is not None and record.get("telegram_id") != telegram_id:
            continue
        if only_errors and not is_error_row(record):
            continue
        if round_filter is not None and record.get("round") != round_filter:
            continue
        rows.append(record)
    return rows


async def query_calls(
    campaign_title: str, telegram_id: Optional[int], only_errors: bool, round_filter: Optional[int]
) -> list[dict]:
    """round_filter=None — все раунды (за всё время); иначе — только звонки этого раунда
    (обычно текущий/последний раунд сценария, см. get_current_round)."""
    async with _lock:
        return await asyncio.to_thread(_query_calls_sync, campaign_title, telegram_id, only_errors, round_filter)


def _call_record(sheet: Worksheet, row: int) -> dict:
    values = [sheet.cell(row=row, column=i + 1).value for i in range(len(CALLS_HEADER))]
    return {**dict(zip(CALLS_HEADER, values)), "_row": row}


def _find_calls_sync(
    campaign_title: Optional[str], telegram_id: Optional[int], text: Optional[str], limit: int
) -> list[dict]:
    sheet = _open_workbook()["Calls"]
    needle = (text or "").strip().lower()
    found = []
    for r in range(sheet.max_row, 1, -1):
        record = _call_record(sheet, r)
        if campaign_title is not None and record.get("campaign") != campaign_title:
            continue
        if telegram_id is not None and record.get("telegram_id") != telegram_id:
            continue
        if needle and not any(needle in str(record.get(k) or "").lower() for k in ("call_id", "sid")):
            continue
        found.append(record)
        if len(found) >= limit:
            break
    return found


async def find_calls(
    campaign_title: Optional[str] = None, telegram_id: Optional[int] = None,
    text: Optional[str] = None, limit: int = 30,
) -> list[dict]:
    """Звонки для просмотра логов, новые сверху. text ищется в call_id и sid.
    У каждой записи есть _row — номер строки листа Calls (см. get_call)."""
    async with _lock:
        return await asyncio.to_thread(_find_calls_sync, campaign_title, telegram_id, text, limit)


def _get_call_sync(row: int) -> Optional[dict]:
    sheet = _open_workbook()["Calls"]
    if not 2 <= row <= sheet.max_row:
        return None
    return _call_record(sheet, row)


async def get_call(row: int) -> Optional[dict]:
    async with _lock:
        return await asyncio.to_thread(_get_call_sync, row)


def _call_users_sync(campaign_title: str) -> list[tuple[int, str]]:
    sheet = _open_workbook()["Calls"]
    users: dict[int, str] = {}
    for r in range(sheet.max_row, 1, -1):
        record = _call_record(sheet, r)
        tid = record.get("telegram_id")
        if record.get("campaign") == campaign_title and tid is not None and tid not in users:
            users[tid] = str(record.get("name") or tid)
    return list(users.items())


async def call_users(campaign_title: str) -> list[tuple[int, str]]:
    """Кто звонил по кампании: (telegram_id, имя), недавние сверху."""
    async with _lock:
        return await asyncio.to_thread(_call_users_sync, campaign_title)


# --- Тест-кейсы ---

def _list_test_plans_sync() -> list[tuple[str, int]]:
    wb = _open_workbook()
    sheet = wb["TestCases"]
    order_seen: list[str] = []
    counts: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        plan = sheet.cell(row=row, column=1).value
        if not plan:
            continue
        if plan not in counts:
            counts[plan] = 0
            order_seen.append(plan)
        counts[plan] += 1
    return [(p, counts[p]) for p in order_seen]


async def list_test_plans() -> list[tuple[str, int]]:
    async with _lock:
        return await asyncio.to_thread(_list_test_plans_sync)


def _get_test_plan_cases_sync(test_plan: str) -> list[dict]:
    wb = _open_workbook()
    sheet = wb["TestCases"]
    rows = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row, column=i + 1).value for i in range(len(TESTCASES_HEADER))]
        record = dict(zip(TESTCASES_HEADER, values))
        if record.get("test_plan") == test_plan:
            rows.append(record)
    rows.sort(key=lambda r: r.get("order") if isinstance(r.get("order"), (int, float)) else 0)
    return rows


async def get_test_plan_cases(test_plan: str) -> list[dict]:
    async with _lock:
        return await asyncio.to_thread(_get_test_plan_cases_sync, test_plan)


def _all_test_cases_sync() -> list[dict]:
    sheet = _open_workbook()["TestCases"]
    rows = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row, column=i + 1).value for i in range(len(TESTCASES_HEADER))]
        if values[0]:
            rows.append(dict(zip(TESTCASES_HEADER, values)))
    return rows


async def all_test_cases() -> list[dict]:
    """Все тест-кейсы всех планов в порядке листа — для выгрузки админу."""
    async with _lock:
        return await asyncio.to_thread(_all_test_cases_sync)


def _replace_test_plan_cases_sync(rows: list[dict]) -> dict[str, int]:
    wb = _open_workbook()
    sheet = wb["TestCases"]
    plans_to_replace = {r["test_plan"] for r in rows}

    kept = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row, column=i + 1).value for i in range(len(TESTCASES_HEADER))]
        record = dict(zip(TESTCASES_HEADER, values))
        if record.get("test_plan") not in plans_to_replace:
            kept.append(values)

    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for values in kept:
        append_row(sheet, values)
    for r in rows:
        append_row(sheet, [r.get(col, "") for col in TESTCASES_HEADER])

    _save_with_retry(wb)

    summary: dict[str, int] = {}
    for r in rows:
        summary[r["test_plan"]] = summary.get(r["test_plan"], 0) + 1
    return summary


async def replace_test_plan_cases(rows: list[dict]) -> dict[str, int]:
    """Заменяет тест-кейсы для планов, встретившихся в rows, остальные планы не трогает."""
    async with _lock:
        return await asyncio.to_thread(_replace_test_plan_cases_sync, rows)


def _find_progress_row(sheet: Worksheet, telegram_id: int, test_plan: str) -> Optional[int]:
    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row=row, column=1).value == telegram_id
            and sheet.cell(row=row, column=2).value == test_plan
        ):
            return row
    return None


def _get_progress_sync(telegram_id: int, test_plan: str) -> int:
    wb = _open_workbook()
    sheet = wb["TestProgress"]
    row = _find_progress_row(sheet, telegram_id, test_plan)
    if row is None:
        return 0
    return sheet.cell(row=row, column=3).value or 0


async def get_progress(telegram_id: int, test_plan: str) -> int:
    async with _lock:
        return await asyncio.to_thread(_get_progress_sync, telegram_id, test_plan)


def _set_progress_sync(telegram_id: int, test_plan: str, index: int) -> None:
    wb = _open_workbook()
    sheet = wb["TestProgress"]
    row = _find_progress_row(sheet, telegram_id, test_plan)
    now = dt.datetime.now().isoformat(timespec="seconds")
    if row is None:
        append_row(sheet, [telegram_id, test_plan, index, now])
    else:
        set_cell(sheet, row, 3, index)
        set_cell(sheet, row, 4, now)
    _save_with_retry(wb)


async def set_progress(telegram_id: int, test_plan: str, index: int) -> None:
    async with _lock:
        await asyncio.to_thread(_set_progress_sync, telegram_id, test_plan, index)


def _delete_rows_where(sheet: Worksheet, predicate) -> int:
    """Удаляет строки (кроме заголовка), для которых predicate(values) истинно.
    Идём снизу вверх, чтобы удаление не сдвигало ещё не просмотренные строки."""
    removed = 0
    for row in range(sheet.max_row, 1, -1):
        values = [sheet.cell(row=row, column=i + 1).value for i in range(sheet.max_column)]
        if predicate(values):
            sheet.delete_rows(row, 1)
            removed += 1
    return removed


def _delete_test_plan_sync(test_plan: str) -> int:
    wb = _open_workbook()
    removed = _delete_rows_where(wb["TestCases"], lambda v: v[0] == test_plan)
    # Прогресс исполнителей по удалённому плану больше не нужен — иначе при повторной
    # загрузке плана с тем же именем они продолжили бы с середины.
    _delete_rows_where(wb["TestProgress"], lambda v: len(v) > 1 and v[1] == test_plan)
    _save_with_retry(wb)
    return removed


async def delete_test_plan(test_plan: str) -> int:
    """Удаляет все тест-кейсы плана и прогресс исполнителей по нему. Возвращает число удалённых кейсов."""
    async with _lock:
        return await asyncio.to_thread(_delete_test_plan_sync, test_plan)
