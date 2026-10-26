import io

from openpyxl import load_workbook

from .campaigns import CAMPAIGNS

REQUIRED_COLUMNS = ["test_plan", "order", "description", "campaign"]
VALID_CAMPAIGN_TITLES = {c.title for c in CAMPAIGNS.values()}


def parse_testcases_file(buf: io.BytesIO) -> tuple[list[dict], list[str]]:
    """Разбирает Excel-файл с тест-кейсами. Возвращает (строки, ошибки).
    Если ошибки не пустые — строки использовать нельзя, показать ошибки админу."""
    try:
        wb = load_workbook(buf, data_only=True)
    except Exception as exc:
        return [], [f"Не удалось открыть файл как Excel: {exc}"]

    sheet = wb.active
    header_row = [
        str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]
    ]
    col_index = {name: idx for idx, name in enumerate(header_row) if name}

    missing = [c for c in REQUIRED_COLUMNS if c not in col_index]
    if missing:
        return [], [f"В файле не хватает колонок: {', '.join(missing)}"]

    rows: list[dict] = []
    errors: list[str] = []

    for r in range(2, sheet.max_row + 1):
        values = [cell.value for cell in sheet[r]]
        if all(v is None or str(v).strip() == "" for v in values):
            continue

        test_plan = str(values[col_index["test_plan"]] or "").strip()
        description = str(values[col_index["description"]] or "").strip()
        campaign = str(values[col_index["campaign"]] or "").strip()
        order_raw = values[col_index["order"]]

        if not test_plan or not description or not campaign:
            errors.append(f"Строка {r}: не заполнены test_plan/description/campaign")
            continue
        if campaign not in VALID_CAMPAIGN_TITLES:
            valid = ", ".join(sorted(VALID_CAMPAIGN_TITLES))
            errors.append(f"Строка {r}: неизвестный сценарий «{campaign}» (допустимо: {valid})")
            continue

        try:
            order = int(order_raw)
        except (TypeError, ValueError):
            order = r

        rows.append({
            "test_plan": test_plan,
            "order": order,
            "description": description,
            "campaign": campaign,
        })

    if not rows and not errors:
        errors.append("В файле нет ни одной строки с данными.")

    return rows, errors
