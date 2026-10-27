"""Админка глазами админа: кампании, списки переменных, тест-планы, пользователи, логи."""
import datetime as dt
import io
import zipfile
from pathlib import Path

from aiogram.exceptions import TelegramBadRequest
from openpyxl import load_workbook

from tests import support
from tests.support import ADMIN_ID, VALID_CURL, documents, joined
from tests.test_storage import xlsx
from bot import campaigns, excel_store, handlers

USER = 100


class CampaignAdminTest(support.BotTestCase):
    async def test_add_campaign_flow(self):
        await self.register(USER)
        await self.press(ADMIN_ID, "adm:camp_add")
        out = await self.send(ADMIN_ID, "curl http://bad")
        self.assertIn("не разобран", joined(out))
        out = await self.send(ADMIN_ID, VALID_CURL)
        self.assertIn("Переменные: brand, color", joined(out))
        self.assertIn("уже есть", joined(await self.send(ADMIN_ID, "Кредиты")))
        out = await self.no_crash(self.send(ADMIN_ID, "  Новая   кампания "))
        campaign = campaigns.by_title("Новая кампания")
        self.assertIsNotNone(campaign)
        card = [r for r in out if r["parse_mode"] == "HTML"]
        self.assertEqual(len(card), 1)
        self.assertIn("PostgreSQL", joined(out))  # БД в тестах не настроена — админ видит предупреждение
        self.assertIn(f"adm:vl_up:{campaign.key}", card[0]["buttons"])

    async def test_cancel_returns_to_menu_and_clears_state(self):
        await self.press(ADMIN_ID, "adm:camp_add")
        await self.press(ADMIN_ID, "adm:cancel")
        out = await self.send(ADMIN_ID, VALID_CURL)
        self.assertNotIn("разобран", joined(out))

    async def test_title_with_html_is_escaped(self):
        await self.no_crash(self.add_campaign_via_bot("<b>Жирная</b> & <script>"))

    async def test_card_for_long_curl(self):
        # Тело курла легко бывает большим (длинные тексты в params) — карточка не должна падать
        curl = VALID_CURL.replace('"lang": "ru"', '"lang": "' + "я" * 5000 + '"')
        await self.press(ADMIN_ID, "adm:camp_add")
        await self.send(ADMIN_ID, curl)
        out = await self.no_crash(self.send(ADMIN_ID, "Длинная"))
        self.assertTrue(any("Длинная" in t for t in support.texts(out)))

    async def test_replace_curl_keeps_identity(self):
        campaigns.all_campaigns()
        await self.press(ADMIN_ID, "adm:camp_repl:credits")
        await self.send(ADMIN_ID, VALID_CURL)
        c = campaigns.get("credits")
        self.assertEqual((c.key, c.title, c.campaign_id), ("credits", "Кредиты", "11111111-2222-3333-4444-555555555555"))

    async def test_delete_campaign_and_stale_buttons(self):
        await self.register(USER)
        await self.press(ADMIN_ID, "adm:camp_delyes:credits")
        self.assertIsNone(campaigns.get("credits"))
        out = await self.no_crash(self.press(USER, "campaign:credits"))
        self.assertEqual(self.triggered, [])
        self.assertTrue(any(r["method"] == "AnswerCallbackQuery" for r in out))
        for data in ("adm:camp:credits", "adm:vl_up:credits", "adm:logc:credits", "adm:logu:credits:0", "newcamp:credits"):
            with self.subTest(data=data):
                self.assertIn("не найдена", joined(await self.no_crash(self.press(ADMIN_ID, data))))

    async def test_call_logs_folder_stays_inside_logs_dir(self):
        # Название кампании задаёт админ и оно становится папкой в call_logs
        logs_dir = Path(support.config.LOGS_DIR).resolve()
        for title in ("..", "../..", "..\\..", " .. ", "Кредиты...", "Тест. "):
            with self.subTest(title=title):
                try:
                    path = handlers._call_log_dir(title, dt.datetime(2026, 10, 27)).resolve()
                except OSError as exc:
                    self.fail(f"{title!r}: папка для логов не создаётся — {exc}")
                self.assertTrue(path.is_relative_to(logs_dir), f"{title!r} -> {path}")


class VarListAdminTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.c = await self.add_campaign_via_bot()

    async def test_upload_download_delete(self):
        await self.press(ADMIN_ID, f"adm:vl_up:{self.c.key}")
        self.assertIn("документ", joined(await self.send(ADMIN_ID, "просто текст")))
        out = await self.send(ADMIN_ID, document=("v.csv", b"brand;color\nkia;\n"))
        self.assertIn("не заполнено color", joined(out))
        out = await self.send(ADMIN_ID, document=("v.csv", "brand;color\nkia;red\nlada;белый\n".encode("utf-8")))
        self.assertIn("2 строк", joined(out))
        self.assertEqual(campaigns.get(self.c.key).var_list.rows, [["kia", "red"], ["lada", "белый"]])

        doc = documents(await self.press(ADMIN_ID, f"adm:vl_down:{self.c.key}"))[0]
        sheet = load_workbook(io.BytesIO(doc["data"])).active
        self.assertEqual([list(r) for r in sheet.iter_rows(values_only=True)],
                         [["brand", "color"], ["kia", "red"], ["lada", "белый"]])

        await self.press(ADMIN_ID, f"adm:vl_del:{self.c.key}")
        self.assertIsNone(campaigns.get(self.c.key).var_list)

    async def test_wrong_extension(self):
        await self.press(ADMIN_ID, f"adm:vl_up:{self.c.key}")
        self.assertIn(".xlsx или .csv", joined(await self.send(ADMIN_ID, document=("v.txt", b"brand,color"))))

    async def test_telegram_refuses_to_give_file(self):
        # Боту Telegram отдаёт файлы до 20 МБ — на больший файл админ должен получить ответ, а не тишину
        await self.press(ADMIN_ID, f"adm:vl_up:{self.c.key}")
        error = TelegramBadRequest(method=None, message="Bad Request: file is too big")
        out = await self.no_crash(self.send(ADMIN_ID, document=("big.xlsx", error)))
        self.assertTrue(support.texts(out))


class TestPlansAdminTest(support.BotTestCase):
    PLAN = [["test_plan", "order", "description", "campaign"], ["План А", 1, "Кейс 1", "Кредиты"],
            ["План А", 2, "Кейс 2", "Узум авто"]]

    async def test_upload_download_delete(self):
        await self.press(ADMIN_ID, "menu:upload_testcases")
        out = await self.send(ADMIN_ID, document=("p.xlsx", xlsx(self.PLAN)))
        self.assertIn("«План А»: 2", joined(out))
        # описание, похожее на формулу, в выгрузке должно остаться текстом
        await excel_store.replace_test_plan_cases([
            {"test_plan": "План А", "order": 1, "description": "Кейс 1", "campaign": "Кредиты"},
            {"test_plan": "План А", "order": 2, "description": "=1+1", "campaign": "Узум авто"},
        ])
        doc = documents(await self.press(ADMIN_ID, "tp_download"))[0]
        sheet = load_workbook(io.BytesIO(doc["data"])).active
        self.assertEqual(sheet.cell(row=3, column=3).data_type, "s")
        self.assertEqual(sheet.cell(row=3, column=3).value, "=1+1")

        await self.press(ADMIN_ID, "tpdel_list")
        await self.press(ADMIN_ID, "tpdel:0")
        await self.press(ADMIN_ID, "tpdel_yes:0")
        self.assertEqual(await excel_store.list_test_plans(), [])

    async def test_unknown_campaign_rejected(self):
        await self.press(ADMIN_ID, "menu:upload_testcases")
        rows = self.PLAN[:2] + [["План Б", 1, "x", "Нет такой"]]
        self.assertIn("неизвестный сценарий", joined(await self.send(ADMIN_ID, document=("p.xlsx", xlsx(rows)))))
        self.assertEqual(await excel_store.list_test_plans(), [])


class UsersAdminTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        await self.register(USER, "Иван", "79991234567")

    async def test_users_button_sends_excel_first(self):
        out = await self.press(ADMIN_ID, "adm:users")
        sends = [r for r in out if r["method"] in ("SendDocument", "SendMessage")]
        self.assertEqual(sends[0]["method"], "SendDocument")
        self.assertTrue(sends[0]["filename"].endswith(".xlsx"))
        self.assertIn("Иван", sends[1]["text"])
        self.assertEqual(documents(await self.press(ADMIN_ID, "adm:users_list")), [])

    async def test_edit_name_and_phone(self):
        await self.press(ADMIN_ID, f"adm:uname:{USER}")
        await self.send(ADMIN_ID, "Пётр")
        await self.press(ADMIN_ID, f"adm:uphone:{USER}")
        self.assertIn("без плюса", joined(await self.send(ADMIN_ID, "+79990000000")))
        await self.send(ADMIN_ID, "998907654321")
        user = (await excel_store.list_users())[0]
        self.assertEqual((user["name"], user["phone"]), ("Пётр", "998907654321"))

    async def test_upload_preview_apply_and_cancel(self):
        users = await excel_store.list_users()
        users[0]["name"] = "Изменённый"
        users.append({"telegram_id": 555, "name": "Новый", "phone": "79990001122", "registered_at": "", "is_admin": False})
        from bot import admin_files
        await self.press(ADMIN_ID, "adm:users_up")
        out = await self.send(ADMIN_ID, document=("u.xlsx", admin_files.build_users_file(users)))
        self.assertIn("Новые (1): Новый", joined(out))
        self.assertIn("Изменятся (1): Изменённый", joined(out))
        await self.press(ADMIN_ID, "adm:cancel")
        self.assertEqual(len(await excel_store.list_users()), 1)

        await self.press(ADMIN_ID, "adm:users_up")
        await self.send(ADMIN_ID, document=("u.xlsx", admin_files.build_users_file(users)))
        await self.press(ADMIN_ID, "adm:users_apply")
        result = {u["telegram_id"]: u for u in await excel_store.list_users()}
        self.assertEqual(result[USER]["name"], "Изменённый")
        self.assertTrue(result[USER]["registered_at"], "дата регистрации не должна теряться")
        self.assertTrue(result[555]["registered_at"])
        self.assertIn("Новый", joined(await self.send(555, "/start")))  # новый пользователь сразу получил доступ

    async def test_user_with_very_long_name(self):
        # Имя задаёт сам пользователь при регистрации (до 4096 символов) — админка не должна падать
        await self.register(300, "Я" * 4090, "79990001122")
        await self.no_crash(self.press(ADMIN_ID, "adm:users"))
        out = await self.no_crash(self.press(ADMIN_ID, "adm:user:300"))
        self.assertTrue(support.texts(out))


class LogsAdminTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        folder = Path(support.config.LOGS_DIR) / "Кредиты" / "2026-10-27"
        folder.mkdir(parents=True)
        for i, tid in enumerate([USER, USER, 300]):
            log = folder / f"call-{i}__iba.txt"
            log.write_text(f"лог {i}", encoding="utf-8")
            await excel_store.add_call_row(
                timestamp=f"2026-10-27T10:0{i}:00", telegram_id=tid, name=f"u{tid}", campaign="Кредиты",
                call_id=f"0101001b-aaaa-{i}", sid=f"SID{i}", log_file=str(log), status="success",
            )

    def zip_names(self, out) -> list[str]:
        names = []
        for doc in documents(out):
            self.assertEqual(doc["timeout"], 600)
            with zipfile.ZipFile(io.BytesIO(doc["data"])) as zf:
                names += zf.namelist()
        return names

    async def test_by_campaign_and_user(self):
        out = await self.press(ADMIN_ID, "adm:logc:credits")
        self.assertIn(f"adm:logu:credits:{USER}", support.all_buttons(out))
        names = self.zip_names(await self.press(ADMIN_ID, f"adm:logu:credits:{USER}"))
        self.assertEqual(sorted(n.split("/")[-1] for n in names if not n.endswith("info.txt")),
                         ["call-0__iba.txt", "call-1__iba.txt"])
        names = self.zip_names(await self.press(ADMIN_ID, "adm:logu:credits:0"))
        self.assertEqual(sum(n.endswith("info.txt") for n in names), 3)

    async def test_search(self):
        await self.press(ADMIN_ID, "adm:logsearch")
        self.assertIn("4 символа", joined(await self.send(ADMIN_ID, "01")))
        self.assertIn("не найдено", joined(await self.send(ADMIN_ID, "zzzzzz")))
        names = self.zip_names(await self.send(ADMIN_ID, "0101001b-aaaa-2"))
        self.assertEqual(sum(n.endswith("info.txt") for n in names), 1)

    async def test_search_too_broad(self):
        for i in range(31):
            await excel_store.add_call_row(campaign="Кредиты", call_id=f"0101001b-bbbb-{i}")
        await self.press(ADMIN_ID, "adm:logsearch")
        out = await self.send(ADMIN_ID, "0101001b")
        self.assertIn("уточните", joined(out))
        self.assertEqual(documents(out), [])

    async def test_campaign_without_calls(self):
        self.assertIn("ещё не было", joined(await self.press(ADMIN_ID, "adm:logc:uzum_auto")))


if __name__ == "__main__":
    import unittest
    unittest.main()
