"""Звонок глазами исполнителя: кампании без переменных, со списком, вручную, тест-кейсы."""
import asyncio
import json

from openpyxl import load_workbook

from tests import support
from tests.support import ADMIN_ID, joined
from bot import campaigns, excel_store, handlers
from bot.api_client import CallResult

USER = 100
PHONE = "998901234567"


class CallFlowTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        await self.register(USER, "Исполнитель", PHONE)
        self.c = await self.add_campaign_via_bot()
        campaigns.set_var_list(self.c.key, campaigns.VariableList(["brand", "color"], [["kia", "red"], ["bmw", "black"]]))

    async def test_campaign_without_variables(self):
        out = await self.press(USER, "campaign:credits")
        url, headers, body = self.triggered[0]
        self.assertEqual(url, "http://campaigns.test/campaigns/dd426938-5558-4e09-9fd1-d039e11e4bca/calls")
        self.assertEqual(body, {"number": PHONE, "params": {}})
        self.assertIn("Курл не отправлен", joined(out))

    async def test_manual_values(self):
        out = await self.press(USER, f"campaign:{self.c.key}")
        self.assertEqual(support.all_buttons(out)[:2], ["var:manual", "var:random"])
        await self.press(USER, "var:manual")
        self.assertIn("пустым", joined(await self.send(USER, "   ")))
        await self.send(USER, "Kia")
        await self.send(USER, "Красный")
        body = self.triggered[0][2]
        self.assertEqual(body, {"number": PHONE, "params": {"lang": "ru", "brand": "Kia", "color": "Красный"}})
        calls = await excel_store.find_calls()
        self.assertEqual(json.loads(calls[0]["params"]), body["params"])

    async def test_list_rotation_only_commits_successful_calls(self):
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:random")
        first = self.triggered[-1][2]["params"]["brand"]
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:random")
        self.assertEqual(self.triggered[-1][2]["params"]["brand"], first, "неудачный звонок не должен тратить строку")

        self.call_result = CallResult(ok=True, sid="S1")
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:random")
        self.assertEqual(self.processed, [self.c.key])
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:random")
        self.assertNotEqual(self.triggered[-1][2]["params"]["brand"], first)

    async def test_no_list_goes_straight_to_manual(self):
        campaigns.set_var_list(self.c.key, None)
        out = await self.press(USER, f"campaign:{self.c.key}")
        self.assertIn("brand (1 из 2)", joined(out))

    async def test_campaign_deleted_during_manual_input(self):
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:manual")
        await self.press(ADMIN_ID, f"adm:camp_delyes:{self.c.key}")
        out = await self.no_crash(self.send(USER, "Kia"))
        self.assertIn("не найдена", joined(out))
        self.assertEqual(self.triggered, [])

    async def test_command_is_not_taken_as_variable_value(self):
        # Исполнитель передумал и набрал команду — звонок со значением "/export" уходить не должен
        await self.press(USER, f"campaign:{self.c.key}")
        await self.press(USER, "var:manual")
        await self.send(USER, "/export")
        await self.send(USER, "Красный")
        self.assertEqual(self.triggered, [])

    async def test_second_call_while_first_running(self):
        gate = asyncio.Event()

        async def slow_trigger(url, headers, body):
            self.triggered.append((url, headers, body))
            await gate.wait()
            return self.call_result

        handlers.trigger_call = slow_trigger
        first = asyncio.create_task(self.press(USER, "campaign:credits"))
        for _ in range(100):
            if self.triggered:
                break
            await asyncio.sleep(0.01)
        out = await self.press(USER, "campaign:credits")
        gate.set()
        await first
        self.assertIn("уже идёт звонок", joined(out))
        self.assertEqual(len(self.triggered), 1)

    async def test_env_admin_without_registration(self):
        out = await self.no_crash(self.press(ADMIN_ID, "campaign:credits"))
        self.assertIn("зарегистрируйтесь", joined(out))
        self.assertEqual(self.triggered, [])


class TestCaseFlowTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        await self.register(USER, "Исполнитель", PHONE)
        self.c = await self.add_campaign_via_bot()
        await excel_store.replace_test_plan_cases([
            {"test_plan": "План", "order": 1, "description": "Со списком", "campaign": self.c.title},
            {"test_plan": "План", "order": 2, "description": "Удалённая", "campaign": "Нет такой"},
            {"test_plan": "План", "order": 3, "description": "Без переменных", "campaign": "Кредиты"},
        ])

    async def test_walk_through_plan(self):
        campaigns.set_var_list(self.c.key, campaigns.VariableList(["brand", "color"], [["kia", "red"]]))
        await self.press(USER, "menu:testcases")
        out = await self.press(USER, "tcplan:0")
        self.assertIn("Тест-кейс 1 из 3", joined(out))
        await self.press(USER, "tc_call")
        self.assertEqual(self.triggered[-1][2]["params"]["brand"], "kia")  # список подставился сам
        out = await self.press(USER, "tc_next")
        self.assertIn("не найден", joined(out))  # кейс с удалённой кампанией пропущен
        self.assertIn("Тест-кейс 3 из 3", joined(out))
        await self.press(USER, "tc_call")
        self.assertEqual(self.triggered[-1][2]["params"], {})
        calls = await excel_store.find_calls()
        self.assertEqual((calls[0]["test_plan"], calls[0]["test_case"]), ("План", "Без переменных"))
        self.assertIn("пройдены", joined(await self.press(USER, "tc_next")))

    async def test_case_without_list_asks_manually(self):
        await self.press(USER, "menu:testcases")
        await self.press(USER, "tcplan:0")
        out = await self.press(USER, "tc_call")
        self.assertIn("brand (1 из 2)", joined(out))
        await self.send(USER, "lada")
        await self.send(USER, "white")
        self.assertEqual(self.triggered[-1][2]["params"]["color"], "white")
        sheet = load_workbook(support.config.EXCEL_PATH)["Calls"]
        self.assertEqual(sheet.cell(row=2, column=excel_store.CALLS_HEADER.index("test_case") + 1).value, "Со списком")


if __name__ == "__main__":
    import unittest
    unittest.main()
