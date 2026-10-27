"""Хранилища и файлы: кампании, ротация переменных, Excel, обмен файлами с админом, архивы."""
import io
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tests import support
from bot import admin_files, campaigns, excel_store, export_service, var_rotation
from bot.curl_parser import parse_curl


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CampaignStoreTest(support.BotTestCase):
    async def test_seed_from_defaults_and_legacy_car_list(self):
        Path(support.config.CAR_MODELS_PATH).write_text(
            "Chevrolet Aveo\nchevrolet  aveo\nBMW X5 M\nOneword\n\n", encoding="utf-8"
        )
        items = campaigns.all_campaigns()
        self.assertEqual(len(items), 6)
        uzum = campaigns.get("uzum_auto")
        self.assertEqual(uzum.variables, ["brand", "model"])
        self.assertEqual(uzum.var_list.rows, [["chevrolet", "aveo"], ["bmw", "x5 m"]])
        self.assertTrue(uzum.url.startswith("http://campaigns.test/campaigns/"))
        self.assertTrue(Path(support.config.CAMPAIGNS_PATH).exists())

    async def test_persist_add_replace_delete(self):
        campaigns.all_campaigns()
        parsed = parse_curl(support.VALID_CURL)
        added = campaigns.add("Новая", parsed)
        campaigns._campaigns = None  # перечитать с диска
        self.assertEqual(campaigns.get(added.key).body, parsed.body)
        campaigns.set_var_list(added.key, campaigns.VariableList(["brand", "color"], [["kia", "red"]]))
        campaigns._campaigns = None
        self.assertIsNone(campaigns.get(added.key).var_list_problem())
        campaigns.delete(added.key)
        campaigns._campaigns = None
        self.assertIsNone(campaigns.get(added.key))

    async def test_var_list_problem(self):
        c = campaigns.add("X", parse_curl(support.VALID_CURL))
        self.assertIn("не загружен", c.var_list_problem())
        campaigns.set_var_list(c.key, campaigns.VariableList(["brand"], [["kia"]]))
        self.assertIn("color", c.var_list_problem())

    async def test_corrupted_file_gives_clear_error(self):
        # Файл правят руками на сервере: битый JSON не должен валить бота непонятной ошибкой
        Path(support.config.CAMPAIGNS_PATH).write_text("{ broken", encoding="utf-8")
        out = await self.no_crash(self.press(support.ADMIN_ID, "adm:camps"))
        self.assertIn("campaigns.json", support.joined(out))


class RotationTest(support.BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.campaign = campaigns.add("Rot", parse_curl(support.VALID_CURL))
        rows = [["kia", "red"], ["bmw", "black"], ["lada", "white"]]
        campaigns.set_var_list(self.campaign.key, campaigns.VariableList(["brand", "color"], rows))

    def take(self, user_id: int) -> dict:
        idx, values = var_rotation.peek_next(user_id, self.campaign)
        var_rotation.commit(user_id, self.campaign, idx)
        return values

    async def test_no_repeats_within_cycle_then_new_cycle(self):
        first = [self.take(1)["brand"] for _ in range(3)]
        self.assertEqual(sorted(first), ["bmw", "kia", "lada"])
        second = [self.take(1)["brand"] for _ in range(3)]
        self.assertEqual(sorted(second), ["bmw", "kia", "lada"])

    async def test_peek_without_commit_repeats_same_row(self):
        self.assertEqual(var_rotation.peek_next(1, self.campaign), var_rotation.peek_next(1, self.campaign))

    async def test_users_independent_and_list_change_resets(self):
        self.take(1)
        self.take(1)
        self.assertEqual(len([self.take(2) for _ in range(3)]), 3)  # у второго свой полный круг
        campaigns.set_var_list(self.campaign.key, campaigns.VariableList(["brand", "color"], [["uaz", "green"]]))
        self.assertEqual(self.take(1), {"brand": "uaz", "color": "green"})


class AdminFilesTest(support.BotTestCase):
    async def test_users_roundtrip_and_validation(self):
        users = [
            {"telegram_id": 1, "name": "=HYPERLINK(\"http://x\")", "phone": "998901234567",
             "registered_at": "2026-01-01T00:00:00", "is_admin": True},
            {"telegram_id": 2, "name": "Б", "phone": "79991234567", "registered_at": "", "is_admin": False},
        ]
        data = admin_files.build_users_file(users)
        sheet = load_workbook(io.BytesIO(data)).active
        self.assertEqual(sheet.cell(row=2, column=2).data_type, "s", "имя-формула должно сохраняться как текст")
        parsed, errors = admin_files.parse_users_file(io.BytesIO(data), "u.xlsx")
        self.assertEqual(errors, [])
        self.assertEqual([(u["telegram_id"], u["name"], u["phone"], u["is_admin"]) for u in parsed],
                         [(1, users[0]["name"], "998901234567", True), (2, "Б", "79991234567", False)])

    async def test_users_numeric_phone_from_excel(self):
        data = xlsx([["telegram_id", "name", "phone"], [123.0, "A", 998901234567.0]])
        parsed, errors = admin_files.parse_users_file(io.BytesIO(data), "u.xlsx")
        self.assertEqual(errors, [])
        self.assertEqual((parsed[0]["telegram_id"], parsed[0]["phone"]), (123, "998901234567"))

    async def test_users_errors(self):
        cases = {
            "колонок": [["id", "name", "phone"], [1, "A", "79991234567"]],
            "числом": [["telegram_id", "name", "phone"], ["abc", "A", "79991234567"]],
            "повторяется": [["telegram_id", "name", "phone"], [1, "A", "79991234567"], [1, "B", "79991234568"]],
            "пустое имя": [["telegram_id", "name", "phone"], [1, "", "79991234567"]],
            "плюса": [["telegram_id", "name", "phone"], [1, "A", "+79991234567"]],
            "ни одного": [["telegram_id", "name", "phone"]],
        }
        for expected, rows in cases.items():
            with self.subTest(expected=expected):
                _, errors = admin_files.parse_users_file(io.BytesIO(xlsx(rows)), "u.xlsx")
                self.assertTrue(any(expected in e for e in errors), errors)

    async def test_var_list_formats(self):
        c = campaigns.add("V", parse_curl(support.VALID_CURL))
        ok = {
            "x.xlsx": xlsx([["color", "brand", "extra", None], ["red", "kia", 1], ["blue", "bmw", 2.0]]),
            "semicolon_cp1251.csv": "brand;color\nкиа;красный\n".encode("cp1251"),
            "comma.csv": "﻿brand,color\nkia,red\n".encode("utf-8"),
        }
        for name, content in ok.items():
            with self.subTest(name=name):
                var_list, errors = admin_files.parse_var_list_file(io.BytesIO(content), name, c)
                self.assertEqual(errors, [])
                self.assertTrue(var_list.rows)

    async def test_var_list_single_column_csv(self):
        c = campaigns.add("One", parse_curl(
            "curl http://h/campaigns/11111111-2222-3333-4444-555555555555/calls -d '{\"number\":\"\",\"p\":\"{{brand}}\"}'"
        ))
        var_list, errors = admin_files.parse_var_list_file(io.BytesIO(b"brand\nkia\nbmw\n"), "one.csv", c)
        self.assertEqual(errors, [])
        self.assertEqual(var_list.rows, [["kia"], ["bmw"]])

    async def test_var_list_errors(self):
        c = campaigns.add("E", parse_curl(support.VALID_CURL))
        cases = {
            "Нет колонок": b"brand\nkia\n",
            "не заполнено": b"brand,color\nkia,\n",
            "без пустых и повторов": b"brand,brand,color\nkia,kia,red\n",
            "нет ни одной строки": b"brand,color\n",
            "пустой": b"",
        }
        for expected, content in cases.items():
            with self.subTest(expected=expected):
                var_list, errors = admin_files.parse_var_list_file(io.BytesIO(content), "e.csv", c)
                self.assertIsNone(var_list)
                self.assertTrue(any(expected in e for e in errors), errors)
        _, errors = admin_files.parse_var_list_file(io.BytesIO(b"not an xlsx"), "e.xlsx", c)
        self.assertTrue(any("Excel" in e for e in errors))


class ExcelStoreTest(support.BotTestCase):
    async def test_old_users_sheet_gets_is_admin_column(self):
        wb = Workbook()
        users = wb.active
        users.title = "Users"
        users.append(["telegram_id", "name", "phone", "registered_at"])
        users.append([5, "Old", "79991234567", "2026-01-01"])
        calls = wb.create_sheet("Calls")
        calls.append(excel_store.CALLS_HEADER[:5])
        wb.save(support.config.EXCEL_PATH)
        self.assertEqual((await excel_store.list_users())[0]["is_admin"], False)
        self.assertTrue(await excel_store.update_user(5, is_admin=True))
        header = [c.value for c in load_workbook(support.config.EXCEL_PATH)["Users"][1]]
        self.assertEqual(header, excel_store.USERS_HEADER)
        self.assertEqual((await excel_store.list_users())[0]["is_admin"], True)

    async def test_user_crud(self):
        await self.register(1, "A", "79991234567")
        await self.register(2, "B", "79991234568")
        self.assertTrue(await excel_store.update_user(1, name="A2", phone="998901234567"))
        self.assertFalse(await excel_store.update_user(999, name="X"))
        self.assertTrue(await excel_store.delete_user(2))
        self.assertFalse(await excel_store.delete_user(2))
        self.assertEqual([(u["telegram_id"], u["name"], u["phone"]) for u in await excel_store.list_users()],
                         [(1, "A2", "998901234567")])

    async def test_calls_lookup(self):
        for i, (campaign, tid) in enumerate([("K1", 1), ("K2", 1), ("K1", 2), ("K1", 1)]):
            await excel_store.add_call_row(timestamp=f"2026-10-0{i + 1}T10:00:00", telegram_id=tid, name=f"u{tid}",
                                           campaign=campaign, call_id=f"call-{i}", sid=f"SID{i}")
        newest = await excel_store.find_calls("K1")
        self.assertEqual([r["call_id"] for r in newest], ["call-3", "call-2", "call-0"])
        self.assertEqual([r["call_id"] for r in await excel_store.find_calls("K1", 1)], ["call-3", "call-0"])
        self.assertEqual([r["call_id"] for r in await excel_store.find_calls(text="sid2")], ["call-2"])
        self.assertEqual(len(await excel_store.find_calls(limit=2)), 2)
        self.assertEqual(await excel_store.call_users("K1"), [(1, "u1"), (2, "u2")])
        row = newest[0]["_row"]
        self.assertEqual((await excel_store.get_call(row))["call_id"], "call-3")
        self.assertIsNone(await excel_store.get_call(1))
        self.assertIsNone(await excel_store.get_call(10_000))

    async def test_comment_formula_stored_as_text(self):
        row = await excel_store.add_call_row(timestamp="t", campaign="K")
        await excel_store.update_call_row(row, comment="=cmd|' /C calc'!A0")
        cell = load_workbook(support.config.EXCEL_PATH)["Calls"].cell(row=row, column=excel_store.CALLS_HEADER.index("comment") + 1)
        self.assertEqual(cell.data_type, "s")


class CallArchivesTest(support.BotTestCase):
    def make_call(self, idx: int, sizes: dict[str, int]) -> dict:
        folder = Path(support.config.LOGS_DIR) / "K" / "2026-10-27"
        folder.mkdir(parents=True, exist_ok=True)
        paths = {}
        for kind, size in sizes.items():
            path = folder / f"call{idx}_{kind}"
            path.write_bytes(b"x" * size)
            paths[kind] = str(path)
        return {
            "_row": idx + 2, "timestamp": f"2026-10-27T10:00:0{idx}", "name": "Ива/н..", "call_id": f"c{idx}",
            "campaign": "K", "log_file": "; ".join(v for k, v in paths.items() if k.startswith("log")),
            "transcript_file": paths.get("tr.txt", ""), "audio_file": paths.get("a.wav", ""),
            "log": "текст лога из таблицы", "transcription": "транскрипция из таблицы",
        }

    def unpack(self, plan) -> list[list[str]]:
        names = []
        for i in range(len(plan.parts)):
            data = plan.build(i)
            self.assertLessEqual(len(data), export_service.ARCHIVE_PART_LIMIT)
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                self.assertIsNone(zf.testzip())
                names.append(zf.namelist())
        return names

    async def test_split_into_parts_under_limit_and_calls_not_torn(self):
        self.addCleanup(setattr, export_service, "ARCHIVE_PART_LIMIT", export_service.ARCHIVE_PART_LIMIT)
        export_service.ARCHIVE_PART_LIMIT = 100_000
        records = [self.make_call(i, {"log1.txt": 20_000, "tr.txt": 5_000, "a.wav": 30_000}) for i in range(5)]
        plan = export_service.plan_calls_archives(records)
        parts = self.unpack(plan)
        self.assertGreater(len(parts), 1)
        all_names = [n for p in parts for n in p]
        self.assertEqual(len(all_names), 5 * 4)  # info + лог + транскрипция + аудио
        for part in parts:  # файлы одного звонка — в одной части
            folders = {n.split("/")[0] for n in part}
            for folder in folders:
                self.assertEqual(sum(n.startswith(folder + "/") for n in all_names),
                                 sum(n.startswith(folder + "/") for n in part))
        for name in all_names:  # имя исполнителя "Ива/н.." не должно давать подпапок и выхода наверх
            self.assertEqual(name.count("/"), 1)
            self.assertNotIn(name.split("/")[0], ("..", "."))

    async def test_oversized_file_listed_in_skipped(self):
        self.addCleanup(setattr, export_service, "ARCHIVE_PART_LIMIT", export_service.ARCHIVE_PART_LIMIT)
        export_service.ARCHIVE_PART_LIMIT = 50_000
        plan = export_service.plan_calls_archives([self.make_call(0, {"log1.txt": 1000, "a.wav": 60_000})])
        names = self.unpack(plan)[0]
        self.assertIn("skipped.txt", names)
        self.assertFalse(any(n.endswith("a.wav") for n in names))

    async def test_fallback_to_table_text_when_files_gone(self):
        record = self.make_call(0, {})
        names = self.unpack(export_service.plan_calls_archives([record]))[0]
        self.assertTrue(any(n.endswith("log_from_table.txt") for n in names))
        self.assertTrue(any(n.endswith("transcript_from_table.txt") for n in names))


if __name__ == "__main__":
    import unittest
    unittest.main()
